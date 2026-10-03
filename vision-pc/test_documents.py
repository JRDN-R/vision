"""Document queue ownership/recovery and real local file conversions (no API calls)."""
import io
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
import time
import unittest
from unittest.mock import patch
import zipfile

import server
import documents
from document_worker import Extractor, preflight_zip
from setup_documents import sample_pdf


class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.config = self.root / 'config.json'
        self.config.write_text(json.dumps({'token':'a'*48, 'documentProcessing':dict(enabled=True,pythonPath=sys.executable,
            profile='Standard',tesseractPath=shutil.which('tesseract'))}))
        server.configure(self.config)
        self.service, self.client = server.documents, server.app.test_client()
        self.service.stop.clear()
        self.headers = {'Authorization':'Bearer '+'a'*48, 'X-Vision-Project-Key':'k'*48}
        self.base='/api/projects/project_documents1'
        self.client.put(self.base,json={'project':{},'revision':0},headers=self.headers)

    def tearDown(self): self.temp.cleanup()

    def upload(self,content=b'Vision evidence 42',name='evidence.txt',receipt='document1',headers=None):
        return self.client.post(self.base+'/documents',data={'requestId':receipt,'file':(io.BytesIO(content),name)},headers=self.headers if headers is None else headers)

    def test_actual_text_job_and_result_cleanup_and_cache(self):
        reply=self.upload();self.assertEqual(reply.status_code,202,reply.json)
        job=reply.json['id'];self.assertTrue(self.service.work_once())
        status=self.client.get(self.base+'/documents/'+job,headers=self.headers).json
        self.assertEqual(status['status'],'complete',status)
        response=self.client.get(self.base+'/documents/'+job+'/result',headers=self.headers)
        data=response.json;response.close()
        self.assertIn('42',data['artifacts'][0]['text'])
        self.assertEqual(len(data['sourceSha256']),64)
        self.assertEqual([p.name for p in (self.service.root/job).iterdir()],['result.json'])
        self.assertEqual(self.upload(receipt='new-request').json['id'],job)

    def test_authentication_ownership_and_receipt_conflict(self):
        self.assertEqual(self.upload(headers={}).status_code,401)
        self.assertEqual(self.upload(headers={**self.headers,'X-Vision-Project-Key':'x'*48}).status_code,403)
        job=self.upload().json['id']
        self.assertEqual(self.upload().json['id'],job)
        self.assertEqual(self.upload(content=b'changed').status_code,409)
        other='/api/projects/project_documents2'
        self.client.put(other,json={'project':{},'revision':0},headers=self.headers)
        for suffix in ('','/result'):
            self.assertEqual(self.client.get(other+'/documents/'+job+suffix,headers=self.headers).status_code,404)
        self.assertEqual(self.client.delete(other+'/documents/'+job,headers=self.headers).status_code,404)
        self.assertEqual(self.client.get(self.base+'/documents/request/document1',headers=self.headers).json['id'],job)

    def test_google_account_ownership(self):
        class Identity:
            project_id='test'
            def verify(self,token): return {'sub':token,'email':token+'@example.test','name':token,'auth_time':int(time.time())}
        server.app.config['FIREBASE_IDENTITY']=Identity()
        a={'Authorization':'Bearer alice'};b={'Authorization':'Bearer bob'}
        account='/api/projects/account_documents1'
        self.assertEqual(self.client.put(account,json={'project':{},'revision':0},headers=a).status_code,200)
        old=self.base;self.base=account
        job=self.upload(headers=a).json['id'];self.base=old
        for endpoint in ('','/'+job,'/'+job+'/result','/request/document1'):
            self.assertEqual(self.client.get(account+'/documents'+endpoint,headers=b).status_code,404)
        self.assertEqual(self.client.delete(account+'/documents/'+job,headers=b).status_code,404)

    def test_restart_cancellation_and_retry_limit(self):
        job=self.upload().json['id']
        self.service.update(job,status='processing')
        self.service.recover();self.assertEqual(self.service.row(job)['status'],'queued')
        self.client.delete(self.base+'/documents/'+job,headers=self.headers)
        self.assertEqual(self.service.row(job)['status'],'cancelled')
        self.assertFalse((self.service.root/job).exists())
        job=self.upload(receipt='another').json['id']
        with server.connect_db() as db: db.execute("UPDATE document_jobs SET attempts=3,status='processing' WHERE id=?",(job,))
        self.service.recover();self.assertEqual(self.service.row(job)['status'],'error')

    def test_limits_and_unsupported_types(self):
        self.assertEqual(self.upload(name='program.exe').status_code,415)
        with patch.object(documents,'UPLOAD_LIMIT',10): self.assertEqual(self.upload().status_code,413)
        with patch.object(documents,'STORAGE_LIMIT',1): self.assertEqual(self.upload().status_code,507)
        self.assertFalse(list(self.service.root.iterdir()))

    def test_real_pdf_word_slides_and_spreadsheet(self):
        from docx import Document
        from pptx import Presentation
        from openpyxl import Workbook
        fixtures={'report.pdf':sample_pdf()}
        doc=Document();doc.add_paragraph('Maintenance evidence 42');table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Part';table.cell(0,1).text='42'
        out=io.BytesIO();doc.save(out);fixtures['report.docx']=out.getvalue()
        ppt=Presentation();slide=ppt.slides.add_slide(ppt.slide_layouts[1]);slide.shapes.title.text='Evidence 42';out=io.BytesIO();ppt.save(out);fixtures['slides.pptx']=out.getvalue()
        book=Workbook();book.active['A1']='Evidence';book.active['B2']=42;book.active['C2']='=B2*2';out=io.BytesIO();book.save(out);fixtures['sheet.xlsx']=out.getvalue()
        for index,(name,content) in enumerate(fixtures.items()):
            job=self.upload(content,name,'fixture'+str(index)).json['id'];self.service.work_once()
            self.assertEqual(self.service.row(job)['status'],'complete',self.service.row(job))
            result=json.loads((self.service.root/job/'result.json').read_text())
            self.assertIn('42',' '.join(a.get('text','') for a in result['artifacts']))
            self.assertTrue(all(a['location'].startswith(name) for a in result['artifacts']))
            if name.endswith('xlsx'):
                self.assertIn('C2',' '.join(a.get('text','') for a in result['artifacts']))

    @unittest.skipUnless(shutil.which('tesseract'),'Tesseract not installed')
    def test_real_image_ocr(self):
        from PIL import Image, ImageDraw
        image=Image.new('RGB',(700,110),'white');ImageDraw.Draw(image).text((20,25),'Vision evidence 42',fill='black',font_size=32)
        out=io.BytesIO();image.save(out,format='PNG')
        job=self.upload(out.getvalue(),'scan.png').json['id'];self.service.work_once()
        self.assertEqual(self.service.row(job)['status'],'complete',self.service.row(job))
        result=json.loads((self.service.root/job/'result.json').read_text())
        self.assertIn('42',' '.join(a.get('text','') for a in result['artifacts']))
        self.assertIn('Tesseract',result['tools'])

    def test_archive_traversal_bombs_and_bounded_supported_members(self):
        for name in ('../secret.txt','/absolute.txt','C:\\private.txt'):
            path=self.root/'unsafe.zip'
            with zipfile.ZipFile(path,'w') as z: z.writestr(name,'42')
            with self.assertRaises(ValueError): preflight_zip(path)
        path=self.root/'bomb.zip'
        with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z: z.writestr('large.txt','a'*2_000_000)
        with self.assertRaises(ValueError): preflight_zip(path)
        out=io.BytesIO()
        with zipfile.ZipFile(out,'w') as z:z.writestr('folder/notes.txt','Evidence 42');z.writestr('program.exe','not executed')
        job=self.upload(out.getvalue(),'bundle.zip').json['id'];self.service.work_once()
        data=json.loads((self.service.root/job/'result.json').read_text())
        self.assertIn('bundle.zip / folder/notes.txt',data['artifacts'][0]['location'])
        self.assertTrue(any('program.exe' in s for s in data['warnings']))

    def test_cancellation_during_completion_cannot_resurrect_result(self):
        job=self.upload().json['id']
        def canceled(value):
            (self.service.root/job/'result.json').write_text('{}')
            self.client.delete(self.base+'/documents/'+job,headers=self.headers)
        with patch.object(self.service,'process',side_effect=canceled):self.service.work_once()
        self.assertEqual(self.service.row(job)['status'],'cancelled')
        self.assertFalse((self.service.root/job).exists())


if __name__=='__main__': unittest.main()
